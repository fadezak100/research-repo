"""Excel workbook explaining E11's avg-dist for ONE query.

Sheets:
  summary      query id, hardness (recall@10 at ef=160), cost, avg-dist,
               pair-distance histogram
  gt_nodes     one row per true neighbor: rank, node id, cosine distance to
               the query, graph degree, hops to its nearest other GT node,
               mean hops to all other GT nodes
  pairs        the board's table: one row per unordered pair (4950 rows) with
               hops in each direction, the min, and the cosine distance
  hop_matrix   100 x 100 symmetrised hop-distance matrix
  paths        actual node-by-node shortest paths for a few example pairs

Usage: .venv/bin/python -m pyhnsw.query_sheet [query_index]
"""

import json
import sys
from collections import deque

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .experiments import RESULTS_DIR, get_ctx
from .graph import DATA_DIR
from .hardness4 import INF, ExactBFS, build_reverse_csr

K_GT = 100
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE8F5")


def shortest_path(adj0, src, dst, max_hops=10):
    """Node-by-node shortest path src -> dst over out-edges (None if too far)."""
    parent = {int(src): None}
    frontier = deque([int(src)])
    for _ in range(max_hops):
        nxt = deque()
        for u in frontier:
            for v in adj0[u]:
                v = int(v)
                if v < 0 or v in parent:
                    continue
                parent[v] = u
                if v == dst:
                    p = [v]
                    while parent[p[-1]] is not None:
                        p.append(parent[p[-1]])
                    return p[::-1]
                nxt.append(v)
        frontier = nxt
    return None


def write_rows(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (
            widths[i - 1] if widths else max(12, len(str(h)) + 2)
        )


def main():
    qi = int(sys.argv[1]) if len(sys.argv) > 1 else 0
    ctx = get_ctx("glove100")
    ds, graph = ctx.ds, ctx.graph
    adj0 = graph.adj0
    q = ds.test[qi]
    gt = ds.ground_truth[qi][:K_GT]
    indptr, _ = build_reverse_csr(adj0, DATA_DIR / "glove100_radj_m16_efc200.npz")

    # directed hop matrix (row i = BFS from GT i), exact
    bfs = ExactBFS(adj0)
    bfs.set_targets(gt)
    Ddir = np.vstack([bfs.hops_from(i) for i in range(K_GT)])
    bfs.clear_targets()
    D = np.minimum(Ddir, Ddir.T)
    iu, ju = np.triu_indices(K_GT, 1)
    pair_hops = D[iu, ju]
    finite = pair_hops[pair_hops < INF]

    e11 = json.loads((RESULTS_DIR / "e11_avgdist_glove100.json").read_text())
    recall = e11["per_query"]["recalls"][qi]
    n_dist = e11["per_query"]["n_dists"][qi]
    d_q = 1.0 - ds.train[gt] @ q
    cos_pair = 1.0 - ds.train[gt] @ ds.train[gt].T

    wb = Workbook()

    # --- summary
    ws = wb.active
    ws.title = "summary"
    hist = {int(h): int(c) for h, c in zip(*np.unique(finite, return_counts=True))}
    rows = [
        ("query index", qi),
        ("dataset", "GloVe-100 (1.18M vectors), HNSW M=16 efC=200, level-0 graph"),
        ("hardness: recall@10 at ef=160", recall),
        ("cost: distance computations at ef=160", n_dist),
        ("ground-truth set", f"exact top-{K_GT} neighbors"),
        ("number of pairs C(100,2)", int(pair_hops.size)),
        ("avg-dist (mean hops over pairs)", float(finite.mean())),
        ("median hops", float(np.median(finite))),
        ("max hops (farthest pair)", int(finite.max())),
        ("unreachable pairs", int((pair_hops == INF).sum())),
        ("", ""),
        ("pair-distance histogram", "number of pairs"),
    ] + [(f"{h} hops", c) for h, c in sorted(hist.items())]
    for r in rows:
        ws.append(r)
    for c in ws["A"]:
        c.font = BOLD
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 60

    # --- gt_nodes
    off = D + np.eye(K_GT, dtype=D.dtype) * INF
    nn_hops = off.min(axis=1)
    mean_hops = np.where(D < INF, D, np.nan).astype(float)
    np.fill_diagonal(mean_hops, np.nan)
    mean_hops = np.nanmean(mean_hops, axis=1)
    outdeg = (adj0[gt] >= 0).sum(axis=1)
    indeg = indptr[gt + 1] - indptr[gt]
    write_rows(
        wb.create_sheet("gt_nodes"),
        ["rank", "node id", "cosine dist to query", "out-degree", "in-degree",
         "hops to nearest other GT", "mean hops to other GT",
         "n GT within 1 hop", "n GT within 2 hops"],
        [(r + 1, int(gt[r]), round(float(d_q[r]), 4), int(outdeg[r]), int(indeg[r]),
          int(nn_hops[r]), round(float(mean_hops[r]), 3),
          int((off[r] <= 1).sum()), int((off[r] <= 2).sum()))
         for r in range(K_GT)],
    )

    # --- pairs
    write_rows(
        wb.create_sheet("pairs"),
        ["rank i", "rank j", "node i", "node j", "hops i->j", "hops j->i",
         "min hops (used)", "cosine dist i-j"],
        [(int(i) + 1, int(j) + 1, int(gt[i]), int(gt[j]),
          int(Ddir[i, j]) if Ddir[i, j] < INF else "unreachable",
          int(Ddir[j, i]) if Ddir[j, i] < INF else "unreachable",
          int(D[i, j]) if D[i, j] < INF else "unreachable",
          round(float(cos_pair[i, j]), 4))
         for i, j in zip(iu, ju)],
    )

    # --- hop_matrix
    ws = wb.create_sheet("hop_matrix")
    ws.append(["rank \\ rank"] + [r + 1 for r in range(K_GT)])
    for r in range(K_GT):
        ws.append([r + 1] + [int(D[r, c]) if D[r, c] < INF else "-" for c in range(K_GT)])
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].font, row[0].fill = BOLD, HEAD
    ws.freeze_panes = "B2"
    for i in range(2, K_GT + 2):
        ws.column_dimensions[get_column_letter(i)].width = 4

    # --- paths: examples at each distance that occurs, plus the farthest pair
    ws = wb.create_sheet("paths")
    ws.append(["example", "hops", "step", "node id", "in GT?", "GT rank",
               "cosine dist to query"])
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
    gt_rank = {int(g): r + 1 for r, g in enumerate(gt)}
    examples = []
    for h in sorted(hist):
        k = np.argmax(pair_hops == h)
        examples.append((f"a pair at {h} hops", int(iu[k]), int(ju[k])))
    for label, i, j in examples:
        a, b = int(gt[i]), int(gt[j])
        p = shortest_path(adj0, a, b) if Ddir[i, j] <= Ddir[j, i] else shortest_path(adj0, b, a)
        for s, node in enumerate(p):
            ws.append([label if s == 0 else "", len(p) - 1 if s == 0 else "", s, node,
                       "GT" if node in gt_rank else "", gt_rank.get(node, ""),
                       round(float(1.0 - ds.train[node] @ q), 4)])
        ws.append([])
    for col, w in zip("ABCDEFG", (22, 6, 6, 12, 8, 9, 20)):
        ws.column_dimensions[col].width = w

    out = RESULTS_DIR / f"e11_query{qi}_avgdist.xlsx"
    wb.save(out)
    print(f"saved {out.relative_to(RESULTS_DIR.parent)}  "
          f"(recall {recall}, avg-dist {finite.mean():.3f}, max {finite.max()})")


if __name__ == "__main__":
    main()

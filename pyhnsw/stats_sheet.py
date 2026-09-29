"""Excel export of the E13 per-query statistics, for reading by hand.

One sheet per dataset with one row per query: the 15 query-time statistics
(computed on the found top-100 of the ef=100 prefix search), their offline
twins on the true top-100, and the targets (recall, oracle ef, cost). A
`columns` sheet explains every column in plain English, and one sheet per
dataset holds the ef sweep (recall and cost at every grid ef). Values only,
no formulas: what is in results/e13_online_stats_{dataset}.json, laid out.

--entry-only writes a slim workbook instead: per query only the hardness
label, recall and the six entry-time statistics (QBAT's five HNSW features
plus ep_nbr_mean), with a `rho` sheet of their Spearman correlations.

Usage:  .venv/bin/python -m pyhnsw.stats_sheet [--entry-only]
Output: results/e13_per_query_stats.xlsx  (or results/e13_entry_features.xlsx)
"""

import argparse
import json

import numpy as np

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .index import RESULTS_DIR
from .online_stats import spearman

DATASETS = ["glove100", "sift1m"]
FONT = "Arial"

# column -> (family, plain-English meaning); order = sheet order
COLUMNS = {
    "query": ("", "row number of the query in the dataset's test set (0-based)"),
    "hard": ("target", "1 if the query is in the hard bin (recall below the dataset's hard line), else 0"),
    "recall": ("target", "recall@100 of the ef=100 prefix search: share of the true top-100 it found"),
    "ef_min_0.8": ("target", "oracle ef: smallest grid ef (100..3200) at which recall reaches 0.8; 6400 = not within the grid"),
    "ef_min_0.9": ("target", "same for recall 0.9"),
    "cost_0.8": ("target", "distance computations of the search at ef_min_0.8"),
    "cost_0.9": ("target", "distance computations of the search at ef_min_0.9"),
    "d10": ("distance", "mean distance from the query to the 10 closest vectors the prefix search found"),
    "d100": ("distance", "mean distance from the query to all 100 found vectors"),
    "contrast10": ("distance", "expected distance to a random base vector divided by d10 (higher = found neighbors stand out more)"),
    "contrast100": ("distance", "same with d100"),
    "spread": ("distance", "mean pairwise distance among the 100 found vectors (higher = scattered)"),
    "edges": ("graph", "number of level-0 links that start and end inside the found top-100 (higher = wired together)"),
    "avgdist": ("graph", "mean shortest-path hops between pairs of the found top-100 (BFS; lab only)"),
    "n_dist": ("effort", "distance computations used by the ef=100 prefix search"),
    "n_hops": ("effort", "nodes expanded by the ef=100 prefix search"),
    "ep_dist": ("entry", "QBAT 1: distance from the query to the layer-0 entry point (where the descent lands)"),
    "ep_nbr_std": ("entry", "QBAT 2: std of the query's distances to the layer-0 entry point's neighbors"),
    "descent_ndist": ("entry", "QBAT 3: distance computations made during the descent through the upper layers"),
    "ep_degree": ("entry", "QBAT 4: number of level-0 links of the layer-0 entry point"),
    "global_ep_dist": ("entry", "QBAT 5: distance from the query to the global entry point (top layer, same node for every query)"),
    "ep_nbr_mean": ("entry", "ours: mean of the query's distances to the layer-0 entry point's neighbors"),
    "d10_gt": ("twin", "d10 computed on the TRUE top-10 instead of the found one"),
    "d100_gt": ("twin", "d100 on the true top-100"),
    "contrast10_gt": ("twin", "contrast10 on the true top-10"),
    "contrast100_gt": ("twin", "contrast100 on the true top-100"),
    "spread_gt": ("twin", "spread of the true top-100"),
    "edges_gt": ("twin", "edges among the true top-100"),
    "avgdist_gt": ("twin", "avg-dist among the true top-100 (E11)"),
}
FAMILY_FILL = {
    "target": "FDE9D9", "distance": "DDEBF7", "graph": "E2EFDA", "effort": "FFF2CC", "entry": "EDE7F6", "twin": "EDEDED",
}
HEAD_FILL = PatternFill("solid", fgColor="2A78D6")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")


def style_header(ws, names):
    for j, name in enumerate(names, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32


def fmt(name, v):
    if v is None:
        return None
    if name in ("hard", "edges", "edges_gt", "n_dist", "n_hops", "ep_degree", "descent_ndist") or name.startswith(("ef_min", "cost")):
        return int(round(v))
    return float(v)


ENTRY = ["ep_dist", "ep_nbr_std", "descent_ndist", "ep_degree", "global_ep_dist", "ep_nbr_mean"]
EASY_LINE = {"glove100": 0.8, "sift1m": 0.95}  # recall at/above which a query is called easy


def entry_only():
    wb = Workbook()
    ws = wb.active
    ws.title = "columns"
    style_header(ws, ["column", "meaning"])
    rows = [("query", COLUMNS["query"][1]),
            ("hardness", "hard = recall below the dataset's hard line (GloVe 0.6, SIFT1M 0.85); easy = recall at or above "
                         "the easy line (GloVe 0.8, SIFT1M 0.95); middle = in between"),
            ("recall", COLUMNS["recall"][1])] + [(f, COLUMNS[f][1]) for f in ENTRY]
    for i, (a, b) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=a).font = Font(name=FONT)
        ws.cell(row=i, column=2, value=b).font = Font(name=FONT)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 16, 120

    rho = wb.create_sheet("rho")
    style_header(rho, ["statistic", "GloVe: rho with recall", "SIFT1M: rho with recall", "how it is computed"])
    for i, f in enumerate(ENTRY, 2):
        rho.cell(row=i, column=1, value=f).font = Font(name=FONT)
    rho.cell(row=2, column=4, value="rank the statistic column and the recall column (smallest = 1, ties get the "
                                    "average position), then correlate the two rank columns; Excel: RANK.AVG on "
                                    "each, then CORREL").font = Font(name=FONT, italic=True)
    for j in (1, 4):
        rho.column_dimensions[get_column_letter(j)].width = 18 if j == 1 else 100
    rho.column_dimensions["B"].width = rho.column_dimensions["C"].width = 22

    for col, ds in ((2, "glove100"), (3, "sift1m")):
        d = json.loads((RESULTS_DIR / f"e13_online_stats_{ds}.json").read_text())
        pq, hb, n = d["per_query"], d["hard_bin"], d["n_queries"]
        rec = np.array(pq["recall"])
        for i, f in enumerate(ENTRY, 2):
            c = rho.cell(row=i, column=col, value=spearman(np.array(pq[f]), rec))
            c.font, c.number_format = Font(name=FONT), "+0.00;-0.00"
        ws = wb.create_sheet(ds)
        names = ["query", "hardness", "recall"] + ENTRY
        style_header(ws, names)
        for i in range(n):
            r = rec[i]
            hard = (r >= hb["lo"]) and ((r <= hb["hi"]) if hb["inclusive"] else (r < hb["hi"]))
            label = "hard" if hard else "easy" if r >= EASY_LINE[ds] else "middle"
            vals = [i, label, float(r)] + [fmt(f, pq[f][i]) for f in ENTRY]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=i + 2, column=j, value=v)
                c.font = Font(name=FONT)
                if label == "hard" and j == 2:
                    c.fill = PatternFill("solid", fgColor=FAMILY_FILL["target"])
                if isinstance(v, float):
                    c.number_format = "0.00" if j == 3 else ("0.000" if abs(v) < 1000 else "#,##0")
        for j in range(1, len(names) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 14
        ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{n + 1}"
    out = RESULTS_DIR / "e13_entry_features.xlsx"
    wb.save(str(out))
    print(f"saved {out.relative_to(RESULTS_DIR.parent)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--entry-only", action="store_true")
    if ap.parse_args().entry_only:
        return entry_only()
    wb = Workbook()
    ws = wb.active
    ws.title = "columns"
    style_header(ws, ["column", "family", "meaning"])
    for i, (name, (fam, meaning)) in enumerate(COLUMNS.items(), 2):
        for j, v in enumerate((name, fam, meaning), 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT)
            if fam:
                c.fill = PatternFill("solid", fgColor=FAMILY_FILL[fam])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 16, 10, 110
    notes = [
        "",
        "Every value comes from results/e13_online_stats_{dataset}.json (E13): 1000 external queries per dataset, "
        "Faiss HNSW M=16 efConstruction=200, prefix search k=100 ef=100.",
        "Distances are the index's own: cosine distance on GloVe-100, squared L2 on SIFT1M. They are not comparable across datasets.",
        "Hard line: GloVe recall < 0.6; SIFT1M recall < 0.85.",
        "Spearman rho between a statistic column and the recall column (rank both, correlate the ranks) gives the "
        "numbers in the E13 tables; Excel has no rank correlation built in, so those are not recomputed here.",
    ]
    for k, line in enumerate(notes, len(COLUMNS) + 3):
        ws.cell(row=k, column=1, value=line).font = Font(name=FONT, italic=True)

    for ds in DATASETS:
        d = json.loads((RESULTS_DIR / f"e13_online_stats_{ds}.json").read_text())
        pq, hb = d["per_query"], d["hard_bin"]
        n = d["n_queries"]
        names = list(COLUMNS)
        ws = wb.create_sheet(ds)
        style_header(ws, names)
        for j, name in enumerate(names, 1):
            fam = COLUMNS[name][0]
            if fam:
                ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor=FAMILY_FILL[fam])
                ws.cell(row=1, column=j).font = Font(name=FONT, bold=True)
        for i in range(n):
            rec = pq["recall"][i]
            hard = (rec >= hb["lo"]) and ((rec <= hb["hi"]) if hb["inclusive"] else (rec < hb["hi"]))
            row = {"query": i, "hard": int(hard)}
            for name in names[2:]:
                row[name] = fmt(name, pq[name][i])
            for j, name in enumerate(names, 1):
                c = ws.cell(row=i + 2, column=j, value=row[name])
                c.font = Font(name=FONT)
                if isinstance(row[name], float):
                    c.number_format = "0.000" if abs(row[name]) < 1000 else "#,##0"
        for j in range(1, len(names) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 13
        ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{n + 1}"

        # the ef sweep: recall and cost at every grid ef
        grid = d["ef_grid"]
        ws2 = wb.create_sheet(f"{ds}_ef_sweep")
        head = ["query"] + [f"recall@ef={e}" for e in grid] + [f"cost@ef={e}" for e in grid]
        style_header(ws2, head)
        for i in range(n):
            vals = [i] + [float(v) for v in pq["recall_grid"][i]] + [int(v) for v in pq["cost_grid"][i]]
            for j, v in enumerate(vals, 1):
                c = ws2.cell(row=i + 2, column=j, value=v)
                c.font = Font(name=FONT)
                if isinstance(v, float):
                    c.number_format = "0.00"
                elif j > 1:
                    c.number_format = "#,##0"
        for j in range(1, len(head) + 1):
            ws2.column_dimensions[get_column_letter(j)].width = 12
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n + 1}"

    out = RESULTS_DIR / "e13_per_query_stats.xlsx"
    wb.save(str(out))
    print(f"saved {out.relative_to(RESULTS_DIR.parent)}")


if __name__ == "__main__":
    main()
